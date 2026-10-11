"""Document extractors — extract structured text from PDF, DOCX, TXT, HTML, CSV, JSON."""

from __future__ import annotations
import json
import logging
from dataclasses import dataclass, field
from pathlib import Path
from typing import List, Optional

logger = logging.getLogger(__name__)


@dataclass
class ExtractedPage:
    """A single page/section from a document."""
    page_number: int
    text: str
    tables: List[str] = field(default_factory=list)
    headings: List[str] = field(default_factory=list)
    metadata: dict = field(default_factory=dict)


@dataclass
class ExtractedDocument:
    """Full extracted document."""
    name: str
    file_type: str
    pages: List[ExtractedPage]
    total_pages: int
    metadata: dict = field(default_factory=dict)


class PDFExtractor:
    """Extract text, tables, and headings from PDF using PyMuPDF."""

    @staticmethod
    def extract(file_path: str) -> ExtractedDocument:
        import fitz  # pymupdf
        doc = fitz.open(file_path)
        pages = []

        for page_num in range(len(doc)):
            page = doc[page_num]
            text = page.get_text("text")

            # Extract tables as text blocks
            tables = []
            table_finder = page.find_tables()
            if table_finder and table_finder.tables:
                for table in table_finder.tables:
                    try:
                        df = table.to_pandas()
                        tables.append(df.to_string(index=False))
                    except Exception:
                        pass

            # Detect headings by font size analysis
            headings = []
            blocks = page.get_text("dict")["blocks"]
            for block in blocks:
                if "lines" in block:
                    for line in block["lines"]:
                        for span in line["spans"]:
                            # Heuristic: larger or bold fonts are headings
                            if span["size"] > 12 or "bold" in span["font"].lower():
                                heading_text = span["text"].strip()
                                if heading_text and len(heading_text) > 2:
                                    headings.append(heading_text)

            pages.append(ExtractedPage(
                page_number=page_num + 1,
                text=text,
                tables=tables,
                headings=headings,
            ))

        doc.close()
        name = Path(file_path).stem
        return ExtractedDocument(
            name=name,
            file_type="pdf",
            pages=pages,
            total_pages=len(pages),
        )


class DOCXExtractor:
    """Extract text and structure from DOCX."""

    @staticmethod
    def extract(file_path: str) -> ExtractedDocument:
        from docx import Document as DocxDocument

        doc = DocxDocument(file_path)
        current_page = 1
        current_text_lines = []
        headings = []
        pages = []

        for para in doc.paragraphs:
            text = para.text.strip()
            if not text:
                current_text_lines.append("")
                continue

            # Detect headings
            if para.style and para.style.name and "Heading" in para.style.name:
                headings.append(text)

            current_text_lines.append(text)

            # Approximate page breaks (DOCX doesn't have real pages)
            if len("\n".join(current_text_lines)) > 3000:
                pages.append(ExtractedPage(
                    page_number=current_page,
                    text="\n".join(current_text_lines),
                    headings=headings.copy(),
                ))
                current_text_lines = []
                headings = []
                current_page += 1

        # Last page
        if current_text_lines:
            pages.append(ExtractedPage(
                page_number=current_page,
                text="\n".join(current_text_lines),
                headings=headings.copy(),
            ))

        # Extract tables
        for i, table in enumerate(doc.tables):
            rows = []
            for row in table.rows:
                cells = [cell.text.strip() for cell in row.cells]
                rows.append(" | ".join(cells))
            table_text = "\n".join(rows)
            # Attach table to the first page (approximation)
            if pages:
                pages[0].tables.append(table_text)

        name = Path(file_path).stem
        return ExtractedDocument(
            name=name,
            file_type="docx",
            pages=pages,
            total_pages=len(pages),
        )


class TXTExtractor:
    """Extract text from plain text files."""

    @staticmethod
    def extract(file_path: str) -> ExtractedDocument:
        text = Path(file_path).read_text(encoding="utf-8", errors="ignore")
        # Split into ~3000 char pages
        chunk_size = 3000
        pages = []
        for i in range(0, len(text), chunk_size):
            page_text = text[i : i + chunk_size]
            pages.append(ExtractedPage(
                page_number=(i // chunk_size) + 1,
                text=page_text,
            ))

        name = Path(file_path).stem
        return ExtractedDocument(
            name=name,
            file_type="txt",
            pages=pages if pages else [ExtractedPage(page_number=1, text="")],
            total_pages=len(pages) or 1,
        )


class HTMLExtractor:
    """Extract text from HTML using BeautifulSoup."""

    @staticmethod
    def extract(file_path: str) -> ExtractedDocument:
        from bs4 import BeautifulSoup

        html = Path(file_path).read_text(encoding="utf-8", errors="ignore")
        soup = BeautifulSoup(html, "html.parser")

        # Remove script and style elements
        for tag in soup(["script", "style", "nav", "footer", "header"]):
            tag.decompose()

        headings = [h.get_text(strip=True) for h in soup.find_all(["h1", "h2", "h3", "h4"])]
        text = soup.get_text(separator="\n", strip=True)

        # Extract tables
        tables = []
        for table in soup.find_all("table"):
            rows = []
            for tr in table.find_all("tr"):
                cells = [td.get_text(strip=True) for td in tr.find_all(["td", "th"])]
                rows.append(" | ".join(cells))
            tables.append("\n".join(rows))

        pages = [ExtractedPage(
            page_number=1,
            text=text,
            tables=tables,
            headings=headings,
        )]

        name = Path(file_path).stem
        return ExtractedDocument(
            name=name,
            file_type="html",
            pages=pages,
            total_pages=1,
        )


class CSVExtractor:
    """Extract data from CSV files."""

    @staticmethod
    def extract(file_path: str) -> ExtractedDocument:
        import pandas as pd

        df = pd.read_csv(file_path)
        # Convert each batch of rows into a page
        rows_per_page = 50
        pages = []
        for i in range(0, len(df), rows_per_page):
            batch = df.iloc[i : i + rows_per_page]
            page_text = batch.to_string(index=False)
            pages.append(ExtractedPage(
                page_number=(i // rows_per_page) + 1,
                text=page_text,
                tables=[page_text],
                headings=list(df.columns) if i == 0 else [],
            ))

        name = Path(file_path).stem
        return ExtractedDocument(
            name=name,
            file_type="csv",
            pages=pages if pages else [ExtractedPage(page_number=1, text="")],
            total_pages=len(pages) or 1,
            metadata={"columns": list(df.columns), "total_rows": len(df)},
        )


class JSONExtractor:
    """Extract data from JSON files."""

    @staticmethod
    def extract(file_path: str) -> ExtractedDocument:
        with open(file_path, "r", encoding="utf-8") as f:
            data = json.load(f)

        text = json.dumps(data, indent=2, ensure_ascii=False)
        # Split into ~3000 char pages
        chunk_size = 3000
        pages = []
        for i in range(0, len(text), chunk_size):
            pages.append(ExtractedPage(
                page_number=(i // chunk_size) + 1,
                text=text[i : i + chunk_size],
            ))

        name = Path(file_path).stem
        return ExtractedDocument(
            name=name,
            file_type="json",
            pages=pages if pages else [ExtractedPage(page_number=1, text="")],
            total_pages=len(pages) or 1,
        )


# ── Extractor registry ──────────────────────────────────────────────────────
EXTRACTORS = {
    ".pdf": PDFExtractor,
    ".docx": DOCXExtractor,
    ".txt": TXTExtractor,
    ".html": HTMLExtractor,
    ".htm": HTMLExtractor,
    ".csv": CSVExtractor,
    ".json": JSONExtractor,
}


def extract_document(file_path: str) -> ExtractedDocument:
    """Extract content from a document based on its file extension."""
    ext = Path(file_path).suffix.lower()
    extractor_cls = EXTRACTORS.get(ext)
    if extractor_cls is None:
        raise ValueError(f"Unsupported file type: {ext}. Supported: {list(EXTRACTORS.keys())}")
    logger.info("Extracting %s with %s", file_path, extractor_cls.__name__)
    return extractor_cls.extract(file_path)
